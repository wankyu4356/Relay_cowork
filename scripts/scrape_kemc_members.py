#!/usr/bin/env python3
"""
한국전기공업협동조합(KEMC) 회원사 목록 크롤링 및 엑셀 저장 스크립트
http://www.kemc.co.kr/member-service/member-list/

사용법:
  1. 기본 (requests 방식): python scrape_kemc_members.py
  2. Selenium 방식:        python scrape_kemc_members.py --selenium
  3. 로컬 HTML 파일:       python scrape_kemc_members.py --local member_list.html
  4. 여러 HTML 파일:       python scrape_kemc_members.py --local page1.html page2.html ...

의존성 설치:
  pip install requests beautifulsoup4 openpyxl lxml
  (Selenium 사용 시): pip install selenium webdriver-manager
"""

import re
import sys
import time
import os
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = "http://www.kemc.co.kr"
MEMBER_LIST_URL = f"{BASE_URL}/member-service/member-list/"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept": (
        "text/html,application/xhtml+xml,application/xml;q=0.9,"
        "image/avif,image/webp,image/apng,*/*;q=0.8"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
    "Accept-Encoding": "gzip, deflate",
    "Connection": "keep-alive",
    "Referer": BASE_URL,
}

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "..", "kemc_members.xlsx")


# ──────────────────────────── HTTP 접근 ────────────────────────────


def create_session():
    """브라우저처럼 동작하는 requests 세션 생성"""
    session = requests.Session()
    session.headers.update(HEADERS)
    try:
        session.get(BASE_URL, timeout=15)
        time.sleep(1)
    except requests.RequestException:
        pass
    return session


def fetch_page(session, url, retries=3):
    """HTTP로 페이지 가져오기 (재시도 포함)"""
    for attempt in range(retries):
        try:
            resp = session.get(url, timeout=15)
            resp.raise_for_status()
            if resp.encoding and resp.encoding.lower() == "iso-8859-1":
                resp.encoding = resp.apparent_encoding
            return resp.text
        except requests.RequestException as e:
            print(f"  [시도 {attempt + 1}/{retries}] 요청 실패: {e}")
            if attempt < retries - 1:
                time.sleep(2 * (attempt + 1))
    return None


# ──────────────────────────── Selenium 접근 ────────────────────────────


def fetch_with_selenium(url, wait_seconds=5):
    """Selenium으로 페이지 가져오기 (실제 브라우저 사용)"""
    try:
        from selenium import webdriver
        from selenium.webdriver.chrome.options import Options
        from selenium.webdriver.chrome.service import Service
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC
    except ImportError:
        print("❌ Selenium이 설치되어 있지 않습니다.")
        print("   설치: pip install selenium webdriver-manager")
        return None

    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_argument(f"user-agent={HEADERS['User-Agent']}")
    options.add_experimental_option("excludeSwitches", ["enable-automation"])

    driver = None
    all_htmls = []
    try:
        # webdriver-manager 자동 설치 시도
        try:
            from webdriver_manager.chrome import ChromeDriverManager
            service = Service(ChromeDriverManager().install())
            driver = webdriver.Chrome(service=service, options=options)
        except ImportError:
            driver = webdriver.Chrome(options=options)

        print(f"  브라우저로 접근 중: {url}")
        driver.get(url)
        time.sleep(wait_seconds)

        # 페이지 로딩 대기
        WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((By.TAG_NAME, "body"))
        )

        page_source = driver.page_source
        all_htmls.append(page_source)
        print(f"  ✅ 페이지 로드 완료 (길이: {len(page_source)})")

        # 페이지네이션 처리: 다음 페이지 클릭
        page_num = 1
        while True:
            page_num += 1
            # 다양한 페이지네이션 셀렉터 시도
            next_page = None
            selectors = [
                f"a[href*='page={page_num}']",
                f"a[href*='pageid={page_num}']",
                f".pagination a:contains('{page_num}')",
                f".kboard-pagination a[href*='{page_num}']",
                f"a.page-numbers:not(.current)",
            ]
            for sel in selectors:
                try:
                    elements = driver.find_elements(By.CSS_SELECTOR, sel)
                    if elements:
                        next_page = elements[0]
                        break
                except Exception:
                    continue

            if not next_page:
                # 숫자 링크로 직접 찾기
                try:
                    all_links = driver.find_elements(By.TAG_NAME, "a")
                    for link in all_links:
                        if link.text.strip() == str(page_num):
                            next_page = link
                            break
                except Exception:
                    pass

            if not next_page:
                break

            try:
                next_page.click()
                time.sleep(2)
                all_htmls.append(driver.page_source)
                print(f"  ✅ 페이지 {page_num} 로드 완료")
            except Exception as e:
                print(f"  페이지 {page_num} 클릭 실패: {e}")
                break

    except Exception as e:
        print(f"  ❌ Selenium 오류: {e}")
    finally:
        if driver:
            driver.quit()

    return all_htmls


# ──────────────────────────── HTML 파싱 ────────────────────────────


def parse_member_list_page(html):
    """회원사 목록 페이지에서 데이터 추출 (다양한 HTML 구조 대응)"""
    soup = BeautifulSoup(html, "lxml")
    members = []

    # 전략 1: 테이블 기반 데이터 추출
    tables = soup.find_all("table")
    for table in tables:
        rows = table.find_all("tr")
        if len(rows) > 1:
            headers = []
            header_row = rows[0]
            for th in header_row.find_all(["th", "td"]):
                headers.append(th.get_text(strip=True))

            if headers:
                for row in rows[1:]:
                    cells = row.find_all(["td", "th"])
                    if cells:
                        member = {}
                        for i, cell in enumerate(cells):
                            key = headers[i] if i < len(headers) else f"컬럼{i+1}"
                            link = cell.find("a")
                            member[key] = cell.get_text(strip=True)
                            if link and link.get("href"):
                                member[f"{key}_링크"] = urljoin(BASE_URL, link["href"])
                        if any(v for k, v in member.items() if not k.endswith("_링크")):
                            members.append(member)

    # 전략 2: KBoard 게시판 리스트 추출
    if not members:
        kboard_list = soup.find("div", class_=re.compile(r"kboard", re.I))
        if kboard_list:
            items = kboard_list.find_all("tr") or kboard_list.find_all(
                "div", class_=re.compile(r"list|item|row", re.I)
            )
            for item in items:
                text = item.get_text(strip=True)
                link = item.find("a")
                if text and link:
                    members.append({
                        "회사명": text,
                        "상세링크": urljoin(BASE_URL, link.get("href", "")),
                    })

    # 전략 3: 일반적인 리스트/카드 구조
    if not members:
        list_containers = soup.find_all(
            "div", class_=re.compile(r"member|company|list|board|entry|post", re.I)
        )
        for container in list_containers:
            links = container.find_all("a")
            for link in links:
                text = link.get_text(strip=True)
                href = link.get("href", "")
                if text and len(text) > 1 and href:
                    members.append({
                        "회사명": text,
                        "상세링크": urljoin(BASE_URL, href),
                    })

    # 전략 4: kboard/member/document 관련 링크 추출
    if not members:
        all_links = soup.find_all("a")
        for link in all_links:
            href = link.get("href", "")
            text = link.get_text(strip=True)
            if text and (
                "kboard_id" in href
                or "member" in href.lower()
                or "document" in href
                or "uid=" in href
            ):
                members.append({
                    "회사명": text,
                    "상세링크": urljoin(BASE_URL, href),
                })

    return members


def find_pagination_urls(html, current_url):
    """페이지네이션 URL 추출"""
    soup = BeautifulSoup(html, "lxml")
    page_urls = set()

    pag_selectors = [
        ("div", re.compile(r"paginat|page-nav|kboard-pagination", re.I)),
        ("nav", re.compile(r"paginat|page|nav-links", re.I)),
        ("ul", re.compile(r"paginat|page", re.I)),
    ]
    for tag, pattern in pag_selectors:
        elements = soup.find_all(tag, class_=pattern)
        for el in elements:
            for link in el.find_all("a"):
                href = link.get("href", "")
                if href:
                    page_urls.add(urljoin(current_url, href))

    return list(page_urls)


def fetch_member_detail(session_or_driver, url, use_selenium=False):
    """개별 회원사 상세 페이지에서 추가 정보 추출"""
    if use_selenium:
        htmls = fetch_with_selenium(url, wait_seconds=3)
        html = htmls[0] if htmls else None
    else:
        html = fetch_page(session_or_driver, url)

    if not html:
        return {}

    soup = BeautifulSoup(html, "lxml")
    detail = {}

    # 테이블 형태 key-value 추출
    tables = soup.find_all("table")
    for table in tables:
        rows = table.find_all("tr")
        for row in rows:
            cells = row.find_all(["th", "td"])
            if len(cells) == 2:
                key = cells[0].get_text(strip=True)
                val = cells[1].get_text(strip=True)
                if key and val:
                    detail[key] = val
            elif len(cells) > 2:
                for i in range(0, len(cells) - 1, 2):
                    key = cells[i].get_text(strip=True)
                    val = cells[i + 1].get_text(strip=True)
                    if key and val:
                        detail[key] = val

    # KBoard 컨텐츠 추출
    if not detail:
        content = soup.find(
            "div", class_=re.compile(r"kboard.*content|entry-content|article", re.I)
        )
        if content:
            detail["내용"] = content.get_text(strip=True)[:500]

    return detail


# ──────────────────────────── 엑셀 저장 ────────────────────────────


def save_to_excel(members, filename):
    """회원사 데이터를 엑셀 파일로 저장"""
    if not members:
        print("저장할 데이터가 없습니다.")
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "KEMC 회원사 목록"

    # 모든 컬럼 키 수집 (순서 유지)
    all_keys = []
    for member in members:
        for key in member.keys():
            if key not in all_keys:
                all_keys.append(key)

    # 스타일 정의
    header_font = Font(name="맑은 고딕", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    data_font = Font(name="맑은 고딕", size=10)
    data_align = Alignment(vertical="center", wrap_text=True)

    # 헤더 작성
    for col_idx, key in enumerate(all_keys, 1):
        cell = ws.cell(row=1, column=col_idx, value=key)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # 데이터 작성
    for row_idx, member in enumerate(members, 2):
        for col_idx, key in enumerate(all_keys, 1):
            value = member.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = data_align
            cell.border = thin_border

    # 열 너비 자동 조정
    for col_idx, key in enumerate(all_keys, 1):
        max_len = len(key) * 2  # 한글은 2배
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    cell_len = len(str(cell.value)) * 1.2
                    max_len = max(max_len, cell_len)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    # 필터 + 고정
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(filename)
    abs_path = os.path.abspath(filename)
    print(f"\n  엑셀 파일 저장 완료: {abs_path}")
    print(f"  총 {len(members)}개 회원사 데이터")
    return True


# ──────────────────────────── 메인 실행 ────────────────────────────


def run_requests_mode():
    """requests 라이브러리로 크롤링"""
    print("[1/5] 세션 생성 중...")
    session = create_session()

    print(f"[2/5] 회원사 목록 페이지 접근: {MEMBER_LIST_URL}")
    html = fetch_page(session, MEMBER_LIST_URL)

    if not html:
        print("\n  페이지 접근 실패! 대안 URL 시도 중...")
        alt_urls = [
            f"{BASE_URL}/?kboard_id=16&mod=list",
            f"{BASE_URL}/?kboard_id=16&mod=list&page=1",
            f"{BASE_URL}/member-service/member-list/?pageid=1",
        ]
        for url in alt_urls:
            print(f"  시도: {url}")
            html = fetch_page(session, url)
            if html:
                break

    if not html:
        return None

    print("[3/5] 회원사 데이터 파싱 중...")
    all_members = parse_member_list_page(html)
    print(f"  현재 페이지에서 {len(all_members)}개 항목 발견")

    print("[4/5] 추가 페이지 확인 중...")
    page_urls = find_pagination_urls(html, MEMBER_LIST_URL)
    visited = {MEMBER_LIST_URL}

    while page_urls:
        url = page_urls.pop(0)
        if url in visited:
            continue
        visited.add(url)
        print(f"  페이지 가져오기: {url}")
        page_html = fetch_page(session, url)
        if page_html:
            page_members = parse_member_list_page(page_html)
            all_members.extend(page_members)
            print(f"    → {len(page_members)}개 항목 추가")
            new_urls = find_pagination_urls(page_html, url)
            for u in new_urls:
                if u not in visited:
                    page_urls.append(u)
            time.sleep(1)

    # 상세 페이지 조회
    members_with_links = [m for m in all_members if m.get("상세링크")]
    if members_with_links and len(members_with_links) <= 500:
        print(f"\n  상세 페이지 {len(members_with_links)}개 조회 중...")
        for i, member in enumerate(members_with_links):
            url = member.get("상세링크", "")
            if url:
                print(f"    [{i+1}/{len(members_with_links)}] {member.get('회사명', '?')}...")
                detail = fetch_member_detail(session, url)
                member.update(detail)
                time.sleep(0.5)

    return all_members


def run_selenium_mode():
    """Selenium(실제 브라우저)으로 크롤링"""
    print("[1/3] 브라우저 시작 중...")
    htmls = fetch_with_selenium(MEMBER_LIST_URL, wait_seconds=5)

    if not htmls:
        print("  브라우저 접근도 실패했습니다.")
        return None

    print(f"[2/3] {len(htmls)}개 페이지 파싱 중...")
    all_members = []
    for i, html in enumerate(htmls):
        page_members = parse_member_list_page(html)
        all_members.extend(page_members)
        print(f"  페이지 {i+1}: {len(page_members)}개 항목")

    return all_members


def run_local_mode(filepaths):
    """로컬 HTML 파일에서 데이터 추출"""
    all_members = []
    for fp in filepaths:
        print(f"  파일 읽는 중: {fp}")
        try:
            # UTF-8 먼저 시도, 실패하면 EUC-KR
            try:
                with open(fp, "r", encoding="utf-8") as f:
                    html = f.read()
            except UnicodeDecodeError:
                with open(fp, "r", encoding="euc-kr") as f:
                    html = f.read()

            page_members = parse_member_list_page(html)
            all_members.extend(page_members)
            print(f"    → {len(page_members)}개 항목 발견")
        except FileNotFoundError:
            print(f"    ❌ 파일을 찾을 수 없습니다: {fp}")

    return all_members


def deduplicate(members):
    """중복 제거"""
    seen = set()
    unique = []
    for m in members:
        key = m.get("회사명", "") or str(m)
        if key not in seen:
            seen.add(key)
            unique.append(m)
    return unique


def main():
    print("=" * 60)
    print("  KEMC 한국전기공업협동조합 회원사 데이터 추출기")
    print("=" * 60)
    print()

    args = sys.argv[1:]
    members = None

    if "--local" in args:
        # 로컬 HTML 파일 모드
        idx = args.index("--local")
        filepaths = args[idx + 1:]
        if not filepaths:
            print("❌ HTML 파일 경로를 지정해주세요.")
            print("   사용법: python scrape_kemc_members.py --local file1.html file2.html ...")
            sys.exit(1)
        members = run_local_mode(filepaths)

    elif "--selenium" in args:
        # Selenium 브라우저 모드
        print("모드: Selenium (실제 브라우저)")
        members = run_selenium_mode()

    else:
        # 기본: requests 모드 → 실패 시 안내
        print("모드: HTTP 요청 (requests)")
        members = run_requests_mode()

    if not members:
        print()
        print("=" * 60)
        print("  ❌ 데이터를 가져올 수 없습니다.")
        print()
        print("  이 사이트는 클라우드/서버 환경에서의 접근을 차단합니다.")
        print("  아래 방법 중 하나로 시도해 주세요:")
        print()
        print("  방법 1: 브라우저에서 HTML 저장 후 실행")
        print("    1) 브라우저에서 아래 URL을 열어주세요:")
        print(f"       {MEMBER_LIST_URL}")
        print("    2) Ctrl+U (페이지 소스 보기) → 전체 선택 → 복사")
        print("    3) member_list.html 파일로 저장")
        print("    4) 실행: python scrape_kemc_members.py --local member_list.html")
        print("    ※ 여러 페이지가 있으면 각각 저장 후:")
        print("       python scrape_kemc_members.py --local page1.html page2.html ...")
        print()
        print("  방법 2: Selenium 사용 (로컬 PC에서)")
        print("    1) pip install selenium webdriver-manager")
        print("    2) python scrape_kemc_members.py --selenium")
        print("=" * 60)
        sys.exit(1)

    members = deduplicate(members)
    print(f"\n  총 {len(members)}개 고유 회원사 데이터")

    print("\n[엑셀 파일 생성 중...]")
    save_to_excel(members, OUTPUT_FILE)
    print("\n  완료!")


if __name__ == "__main__":
    main()
